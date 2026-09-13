"""
Tests de la lógica de src/organizador.py con PDFs generados al vuelo.
Se ejecuta:  python3.12 tests_organizador.py
"""
import sys
import tempfile
from pathlib import Path

try:
    import pymupdf as fitz
except ImportError:
    import fitz

sys.path.insert(0, str(Path(__file__).parent))
from src import anonimizador, compresor, conversor_docx, organizador  # noqa: E402

FALLOS = []


def check(nombre, condicion, detalle=""):
    if condicion:
        print(f"  ✓ {nombre}")
    else:
        print(f"  ✗ {nombre}  {detalle}")
        FALLOS.append(nombre)


def crear_pdf(ruta, textos, con_metadatos=False):
    """Crea un PDF con una página por cada texto (el texto identifica la página)."""
    doc = fitz.open()
    for t in textos:
        pagina = doc.new_page(width=595, height=842)
        pagina.insert_text((72, 100), t, fontsize=24)
    if con_metadatos:
        doc.set_metadata({"author": "Autor Prueba", "title": "Titulo Prueba"})
    doc.save(ruta)
    doc.close()


def textos_de(ruta):
    """Devuelve el texto (limpio) de cada página."""
    doc = fitz.open(ruta)
    try:
        return [doc[i].get_text().strip() for i in range(len(doc))]
    finally:
        doc.close()


def rotaciones_de(ruta):
    doc = fitz.open(ruta)
    try:
        return [doc[i].rotation for i in range(len(doc))]
    finally:
        doc.close()


with tempfile.TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    pdf_a = tmp / "A.pdf"
    pdf_b = tmp / "B.pdf"
    crear_pdf(pdf_a, ["A1", "A2", "A3"], con_metadatos=True)
    crear_pdf(pdf_b, ["B1", "B2"])

    # ---------------- numero_paginas / abrir_pdf ----------------
    print("abrir_pdf / numero_paginas:")
    check("cuenta páginas", organizador.numero_paginas(pdf_a) == 3)
    try:
        organizador.abrir_pdf(tmp / "no_existe.pdf")
        check("error si no existe", False)
    except ValueError:
        check("error si no existe", True)

    # ---------------- UNIR ----------------
    print("unir_pdfs:")
    salida_union = tmp / "union.pdf"
    eventos = []
    res = organizador.unir_pdfs(
        [pdf_a, pdf_b], salida_union,
        organizador.OpcionesUnion(),
        progreso_callback=lambda t, i, n: eventos.append((t, i, n)),
    )
    check("páginas 3+2=5", res.num_paginas == 5)
    check("orden correcto", textos_de(salida_union) == ["A1", "A2", "A3", "B1", "B2"])
    check("callback progreso", len(eventos) == 2)
    check("tamaño > 0", res.tamano_final > 0)

    # orden inverso
    salida_union2 = tmp / "union2.pdf"
    organizador.unir_pdfs([pdf_b, pdf_a], salida_union2)
    check("orden inverso", textos_de(salida_union2) == ["B1", "B2", "A1", "A2", "A3"])

    # mínimo 2 PDFs
    try:
        organizador.unir_pdfs([pdf_a], tmp / "x.pdf")
        check("error con 1 solo PDF", False)
    except ValueError:
        check("error con 1 solo PDF", True)

    # unión con anonimización
    salida_union3 = tmp / "union3.pdf"
    organizador.unir_pdfs(
        [pdf_a, pdf_b], salida_union3,
        organizador.OpcionesUnion(anonimizar=True),
    )
    meta = fitz.open(salida_union3).metadata
    check("anonimizado (autor vacío)", not meta.get("author"))

    # ---------------- PARSEAR RANGOS ----------------
    print("parsear_rangos:")
    check("rangos válidos",
          organizador.parsear_rangos("1-2, 3, 2-3", 3) == [(1, 2), (3, 3), (2, 3)])
    check("espacios raros",
          organizador.parsear_rangos(" 1 - 2 ,3 ", 3) == [(1, 2), (3, 3)])
    for texto, motivo in [
        ("", "vacío"), ("a-b", "letras"), ("0-2", "página 0"),
        ("3-1", "invertido"), ("2-9", "fuera de rango"), (",,,", "solo comas"),
    ]:
        try:
            organizador.parsear_rangos(texto, 3)
            check(f"rechaza {motivo}", False)
        except ValueError as e:
            check(f"rechaza {motivo}", True, str(e))

    # ---------------- DIVIDIR ----------------
    print("dividir_pdf:")
    pdf_c = tmp / "C.pdf"
    crear_pdf(pdf_c, [f"C{i}" for i in range(1, 8)])   # 7 páginas

    # por rangos
    carpeta_r = tmp / "rangos"
    res = organizador.dividir_pdf(
        pdf_c, carpeta_r,
        organizador.OpcionesDivision(modo="rangos", rangos_texto="1-3, 5, 6-7"),
    )
    check("rangos: 3 PDFs", len(res.pdfs_generados) == 3)
    check("rangos: contenido 1-3",
          textos_de(res.pdfs_generados[0]) == ["C1", "C2", "C3"])
    check("rangos: contenido 5", textos_de(res.pdfs_generados[1]) == ["C5"])
    check("rangos: contenido 6-7",
          textos_de(res.pdfs_generados[2]) == ["C6", "C7"])
    check("rangos: nombres",
          res.pdfs_generados[0].name == "C_pags_1-3.pdf"
          and res.pdfs_generados[1].name == "C_pag_5.pdf")

    # cada N
    carpeta_n = tmp / "cada_n"
    res = organizador.dividir_pdf(
        pdf_c, carpeta_n, organizador.OpcionesDivision(modo="cada_n", cada_n=3),
    )
    check("cada 3: 3 PDFs (3+3+1)", len(res.pdfs_generados) == 3)
    check("cada 3: último con 1 página",
          textos_de(res.pdfs_generados[-1]) == ["C7"])

    # página a página
    carpeta_p = tmp / "paginas"
    res = organizador.dividir_pdf(
        pdf_c, carpeta_p, organizador.OpcionesDivision(modo="paginas"),
    )
    check("página a página: 7 PDFs", len(res.pdfs_generados) == 7)
    check("página a página: contenido",
          [textos_de(p)[0] for p in res.pdfs_generados] ==
          [f"C{i}" for i in range(1, 8)])

    # no sobrescribe si se repite (ruta_salida_unica)
    res2 = organizador.dividir_pdf(
        pdf_c, carpeta_p, organizador.OpcionesDivision(modo="paginas"),
    )
    check("no sobrescribe (nombres únicos)",
          all(p.exists() for p in res.pdfs_generados + res2.pdfs_generados)
          and len(set(res.pdfs_generados + res2.pdfs_generados)) == 14)

    # ---------------- RECOMPONER ----------------
    print("recomponer_pdf:")
    PE = organizador.PaginaEditada

    # reordenar: invertir
    salida_re = tmp / "reordenado.pdf"
    res = organizador.recomponer_pdf(
        pdf_a, [PE(2), PE(1), PE(0)], salida_re,
    )
    check("orden invertido", textos_de(salida_re) == ["A3", "A2", "A1"])

    # eliminar (solo 2 de 3) + rotar
    salida_rot = tmp / "rotado.pdf"
    organizador.recomponer_pdf(
        pdf_a, [PE(0, 90), PE(2, 180)], salida_rot,
    )
    check("eliminación (2 páginas)", textos_de(salida_rot) == ["A1", "A3"])
    check("rotaciones 90/180", rotaciones_de(salida_rot) == [90, 180])

    # extraer una sola página
    salida_ext = tmp / "extracto.pdf"
    res = organizador.recomponer_pdf(pdf_a, [PE(1)], salida_ext)
    check("extracción página 2", textos_de(salida_ext) == ["A2"])
    check("resultado num_paginas", res.num_paginas == 1)

    # duplicar una página (permitido por diseño)
    salida_dup = tmp / "dup.pdf"
    organizador.recomponer_pdf(pdf_a, [PE(0), PE(0)], salida_dup)
    check("duplicado", textos_de(salida_dup) == ["A1", "A1"])

    # errores
    try:
        organizador.recomponer_pdf(pdf_a, [], tmp / "y.pdf")
        check("error lista vacía", False)
    except ValueError:
        check("error lista vacía", True)
    try:
        organizador.recomponer_pdf(pdf_a, [PE(99)], tmp / "z.pdf")
        check("error página fuera de rango", False)
    except ValueError:
        check("error página fuera de rango", True)

    # rotación 360 = sin rotación efectiva
    salida_360 = tmp / "r360.pdf"
    organizador.recomponer_pdf(pdf_a, [PE(0, 360 % 360)], salida_360)
    check("rotación 0/360", rotaciones_de(salida_360) == [0])

    # ---------------- PDFs PROTEGIDOS CON CONTRASEÑA ----------------
    print("PDFs cifrados:")
    pdf_cifrado = tmp / "cifrado.pdf"
    d = fitz.open()
    d.new_page().insert_text((72, 100), "secreto", fontsize=20)
    d.save(pdf_cifrado, encryption=fitz.PDF_ENCRYPT_AES_256,
           user_pw="clave", owner_pw="clave")
    d.close()
    try:
        organizador.abrir_pdf(pdf_cifrado)
        check("organizador rechaza cifrado", False)
    except ValueError as e:
        check("organizador rechaza cifrado", "contraseña" in str(e))
    try:
        compresor.comprimir_pdf(
            pdf_cifrado, tmp / "no.pdf", compresor.OpcionesCompresion(),
        )
        check("compresor rechaza cifrado", False)
    except ValueError as e:
        check("compresor rechaza cifrado", "contraseña" in str(e))

    # ---------------- ANONIMIZADOR: limpiar_exif ----------------
    print("limpiar_exif:")
    from PIL import Image as _Img
    original = _Img.new("RGB", (64, 40))
    original.putdata([(x % 256, (x * 3) % 256, (x * 7) % 256)
                      for x in range(64 * 40)])
    original.info["exif"] = b"datos falsos"
    limpia = anonimizador.limpiar_exif(original)
    check("píxeles idénticos", list(limpia.getdata()) == list(original.getdata()))
    check("sin metadatos info", not limpia.info)
    check("mismo modo y tamaño",
          limpia.mode == original.mode and limpia.size == original.size)

    # ---------------- MINIATURAS ----------------
    print("render_miniatura_png:")
    doc = fitz.open(pdf_a)
    png = organizador.render_miniatura_png(doc, 0, 0, max_px=200)
    check("PNG válido", png[:8] == b"\x89PNG\r\n\x1a\n")
    import io
    from PIL import Image
    img = Image.open(io.BytesIO(png))
    check("lado mayor ≤ 200", max(img.size) <= 201, str(img.size))
    w0, h0 = img.size
    png90 = organizador.render_miniatura_png(doc, 0, 90, max_px=200)
    img90 = Image.open(io.BytesIO(png90))
    check("giro 90° intercambia lados", img90.size == (h0, w0),
          f"{img.size} vs {img90.size}")
    doc.close()

    # ---------------- PDF A WORD (pdf2docx) ----------------
    print("convertir_pdf_a_docx:")
    from docx import Document as _Docx

    # PDF con título, párrafo, tabla con rejilla y 2 páginas extra
    pdf_word = tmp / "informe.pdf"
    dw = fitz.open()
    pg = dw.new_page(width=595, height=842)
    pg.insert_text((72, 80), "Informe de prueba", fontsize=22)
    pg.insert_text((72, 130),
                   "Parrafo normal del informe clinico con texto corrido.",
                   fontsize=11)
    tx, ty, tw, th = 72, 180, 200, 60
    for i in range(3):
        pg.draw_line((tx, ty + i * th / 2), (tx + tw, ty + i * th / 2))
        pg.draw_line((tx + i * tw / 2, ty), (tx + i * tw / 2, ty + th))
    for cx, cy, t in [(80, 200, "A1"), (180, 200, "B1"),
                      (80, 230, "A2"), (180, 230, "B2")]:
        pg.insert_text((cx, cy), t, fontsize=10)
    for n in (2, 3):
        pgn = dw.new_page(width=595, height=842)
        pgn.insert_text((72, 100), f"PAGINA{n}", fontsize=14)
    dw.save(pdf_word)
    dw.close()

    salida_docx = tmp / "informe.docx"
    res = conversor_docx.convertir_pdf_a_docx(
        pdf_word, salida_docx, conversor_docx.OpcionesPdfADocx(),
    )
    check("docx generado", salida_docx.exists() and res.tamano_final > 0)
    check("páginas convertidas = 3", res.num_paginas == 3)
    check("sin aviso de escaneado", res.aviso is None)

    d = _Docx(str(salida_docx))
    texto_total = " ".join(p.text for p in d.paragraphs)
    celdas = [c.text.strip() for t in d.tables
              for fila in t.rows for c in fila.cells]
    check("texto convertido", "Informe de prueba" in texto_total
          and "texto corrido" in texto_total)
    check("tabla detectada", celdas == ["A1", "B1", "A2", "B2"], str(celdas))
    check("autor = ARAEMKA PDF", d.core_properties.author == "ARAEMKA PDF")
    check("título = nombre del PDF", d.core_properties.title == "informe")

    # rango de páginas: solo la 2
    salida_p2 = tmp / "informe_p2.docx"
    res = conversor_docx.convertir_pdf_a_docx(
        pdf_word, salida_p2,
        conversor_docx.OpcionesPdfADocx(rangos_texto="2"),
    )
    t2 = " ".join(p.text for p in _Docx(str(salida_p2)).paragraphs)
    check("rango: solo página 2",
          "PAGINA2" in t2 and "PAGINA3" not in t2
          and "Informe de prueba" not in t2)
    check("rango: num_paginas = 1", res.num_paginas == 1)

    # rango fuera de límites
    try:
        conversor_docx.convertir_pdf_a_docx(
            pdf_word, tmp / "x.docx",
            conversor_docx.OpcionesPdfADocx(rangos_texto="9"),
        )
        check("rango fuera de límites rechazado", False)
    except ValueError:
        check("rango fuera de límites rechazado", True)

    # anonimización del Word
    salida_anon = tmp / "informe_anon.docx"
    conversor_docx.convertir_pdf_a_docx(
        pdf_word, salida_anon,
        conversor_docx.OpcionesPdfADocx(anonimizar=True),
    )
    cp = _Docx(str(salida_anon)).core_properties
    check("docx anonimizado", cp.author == "" and cp.title == "")

    # PDF "escaneado" (solo imagen) → aviso
    import io as _io
    from PIL import Image as _Img2
    img = _Img2.new("RGB", (300, 400), (240, 240, 240))
    buf = _io.BytesIO()
    img.save(buf, format="JPEG")
    pdf_escaneado = tmp / "escaneado.pdf"
    de = fitz.open()
    pe = de.new_page(width=595, height=842)
    pe.insert_image(fitz.Rect(50, 50, 545, 792), stream=buf.getvalue())
    de.save(pdf_escaneado)
    de.close()

    salida_esc = tmp / "escaneado.docx"
    res = conversor_docx.convertir_pdf_a_docx(pdf_escaneado, salida_esc)
    check("escaneado: docx generado igualmente", salida_esc.exists())
    check("escaneado: aviso presente",
          res.aviso is not None and "OCR" in res.aviso)

    # PDF cifrado rechazado también aquí
    try:
        conversor_docx.convertir_pdf_a_docx(pdf_cifrado, tmp / "no2.docx")
        check("cifrado rechazado en PDF→Word", False)
    except ValueError as e:
        check("cifrado rechazado en PDF→Word", "contraseña" in str(e))

print()
if FALLOS:
    print(f"RESULTADO: {len(FALLOS)} FALLOS -> {FALLOS}")
    sys.exit(1)
print("RESULTADO: TODOS LOS TESTS OK")
